import os
import re
import logging
import threading
from datetime import datetime

from django.utils import timezone
from docx import Document as DocxDocument
from PyPDF2 import PdfReader

from .models import Document, DocumentChunk

logger = logging.getLogger(__name__)

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150
MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB

ALLOWED_EXTENSIONS = {'.docx', '.pdf'}


def validate_upload(file):
    errors = []

    if file.size > MAX_FILE_SIZE:
        errors.append(f'File too large. Maximum size is {MAX_FILE_SIZE // (1024*1024)}MB.')

    ext = os.path.splitext(file.name)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        errors.append(f'Invalid file type: {ext}. Allowed: {", ".join(ALLOWED_EXTENSIONS)}')

    safe_name = re.sub(r'[^\w\s\-.]', '', file.name)
    if not safe_name:
        errors.append('Invalid filename.')

    return errors


def extract_text_from_docx(file_path):
    doc = DocxDocument(file_path)
    pages = []
    current_page = []
    page_num = 1

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style and para.style.name and 'Heading' in para.style.name:
            if current_page:
                pages.append({'page': page_num, 'text': '\n'.join(current_page)})
                current_page = []
                page_num += 1

        current_page.append(text)

    if current_page:
        pages.append({'page': page_num, 'text': '\n'.join(current_page)})

    if not pages:
        full_text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        pages.append({'page': 1, 'text': full_text})

    return pages


def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    pages = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages.append({'page': i + 1, 'text': text})

    return pages


def extract_text(file_path, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    return extract_text_from_docx(file_path)


def create_chunks(pages, document, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    DocumentChunk.objects.filter(document=document).delete()

    chunks = []
    chunk_index = 0

    for page_data in pages:
        page_num = page_data['page']
        text = page_data['text']

        if not text.strip():
            continue

        start = 0
        while start < len(text):
            end = start + chunk_size
            chunk_text = text[start:end].strip()

            if chunk_text:
                chunks.append(DocumentChunk(
                    document=document,
                    content=chunk_text,
                    chunk_index=chunk_index,
                    page_number=page_num,
                ))
                chunk_index += 1

            start += chunk_size - overlap

    if chunks:
        DocumentChunk.objects.bulk_create(chunks)

    return len(chunks)


def process_document(document_id):
    try:
        doc = Document.objects.get(id=document_id)
    except Document.DoesNotExist:
        logger.error(f"Document {document_id} not found")
        return False

    doc.status = 'processing'
    doc.error_message = None
    doc.save(update_fields=['status', 'error_message'])

    try:
        pages = extract_text(doc.file.path, doc.file.name)

        if not pages:
            raise ValueError("No text could be extracted from the document")

        full_text = '\n\n'.join(p['text'] for p in pages)
        doc.full_text = full_text
        doc.save(update_fields=['full_text'])

        chunk_count = create_chunks(pages, doc)

        doc.status = 'ready'
        doc.processed_at = timezone.now()
        doc.save(update_fields=['status', 'processed_at'])

        try:
            from .indexing.index_sync import upsert_document_chunks
            upsert_document_chunks(doc.id)
        except Exception as e:
            logger.warning(f"Vector index upsert failed for doc {doc.id}: {e}")

        logger.info(f"Document {document_id} processed: {chunk_count} chunks created")
        return True

    except Exception as e:
        error_msg = str(e)
        logger.error(f"Document {document_id} processing failed: {error_msg}")
        doc.status = 'failed'
        doc.error_message = error_msg
        doc.save(update_fields=['status', 'error_message'])
        return False


def process_document_async(document_id):
    thread = threading.Thread(target=process_document, args=(document_id,))
    thread.daemon = True
    thread.start()
    return thread
