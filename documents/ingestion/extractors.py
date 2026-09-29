import os
import logging

from docx import Document as DocxDocument
from PyPDF2 import PdfReader

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    pages = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages.append({'page': i + 1, 'text': text})

    return pages


def extract_text_from_docx(file_path):
    doc = DocxDocument(file_path)
    pages = []
    current_page = []
    page_num = 1

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style and para.style.name and 'Heading' in para.style.name:
            if current_page:
                pages.append({'page': page_num, 'text': '\n'.join(current_page)})
                current_page = []
                page_num += 1

        current_page.append(text)

    if current_page:
        pages.append({'page': page_num, 'text': '\n'.join(current_page)})

    if not pages:
        full_text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        pages.append({'page': 1, 'text': full_text})

    return pages


def extract_text(file_path, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")